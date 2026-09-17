"""Generates CyberHealth_Sheet_v2.xlsx: the restructured, AppSheet-style database.

Run: python3 scripts/make_template_v2.py
"""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.comments import Comment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "CyberHealth_Sheet_v2.xlsx")
MAX_ROW = 1000
PHOTO_FOLDER_ID = "1FJxIDbUBCJR8aRW4gbag4I_mdtxgUiuA"

FILL_KEY = PatternFill("solid", fgColor="4B2E83")
FILL_REQ = PatternFill("solid", fgColor="C00000")
FILL_OPT = PatternFill("solid", fgColor="1F4E78")
FILL_APP = PatternFill("solid", fgColor="7F7F7F")
WHITE = Font(bold=True, color="FFFFFF")
EXAMPLE = Font(italic=True, color="808080")

LISTS = {
    "Role": ["Primary", "Family"],
    "YesNo": ["TRUE", "FALSE"],
    "Section": ["Medicines", "Care team"],
    "Access": ["View", "Edit"],
    "Form": ["Tablet", "Capsule", "Liquid", "Injection", "Inhaler", "Cream", "Drops", "Patch", "Other"],
    "Unit": ["tablet", "capsule", "ml", "unit", "injection", "puff", "drop", "patch", "other"],
    "MealTiming": ["Before meal", "After meal", "With meal", "Any time"],
    "Frequency": ["Daily", "Every N days", "Weekdays", "As needed"],
    "TimeOfDay": ["Morning", "Noon", "Evening", "Bedtime"],
    "Weekday": ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"],
    "Status": ["Active", "Stopped"],
    "ChangeType": ["Started", "Dose changed", "Schedule changed", "Stopped", "Restarted", "Corrected"],
    "DoseStatus": ["Taken", "Skipped"],
    "BloodType": ["A+", "A-", "B+", "B-", "AB+", "AB-", "O+", "O-", "Unknown"],
}

# Column: (name, kind, who, required, type label, example, description, validation)
#   kind: "key" | "col"; who: "you" | "app"
#   validation: None | ("list", ListName) | ("ref", "Table") | ("number",)
def c(name, who, req, typ, example, desc, val=None, kind="col"):
    return (name, kind, who, req, typ, example, desc, val)

def key(name, example, desc):
    return c(name, "you", True, "Key (text, unique)", example, desc, None, "key")

AUDIT = [
    c("created_at", "app", False, "DateTime", "", "When this row was created. Filled by the app."),
    c("created_by", "app", False, "Ref → Users", "", "Who created it. Filled by the app."),
    c("updated_at", "app", False, "DateTime", "", "Last change. Filled by the app."),
    c("updated_by", "app", False, "Ref → Users", "", "Who made the last change. Filled by the app."),
]

TABLES = {
    # ---------- People and access ----------
    "Users": {
        "purpose": "Everyone who logs in. Added by the family admin in the Sheet, never from the app.",
        "cols": [
            key("user_id", "U01", "Short ID for the person, e.g. U01. Never change it once used."),
            c("display_name", "you", True, "Text", "Dad", "Name shown in the app and typed at log in. Must be unique."),
            c("role", "you", True, "Enum", "Primary", "Primary = the person the app opens on by default (Dad). Everyone else is Family.", ("list", "Role")),
            c("password_hash", "app", False, "Text", "", "Filled by the app. Never type a password here. Copy it across from the old Sheet to keep someone's password."),
            c("reset_code", "you", False, "Text (6 digits)", "", "One-time code for setting a password. The app clears it after use."),
            c("active", "you", True, "Yes/No", "TRUE", "FALSE stops this person logging in.", ("list", "YesNo")),
            c("created_at", "app", False, "DateTime", "", "Filled by the app."),
        ],
        "examples": [
            ["U01", "Dad", "Primary", "", "", "TRUE", ""],
            ["U02", "Ooa", "Family", "", "", "TRUE", ""],
        ],
    },
    "Sharing": {
        "purpose": "Who may see or edit someone's records. Only the owner can change their own rows from the app.",
        "cols": [
            key("sharing_id", "SH01", "Unique ID for this rule."),
            c("owner_user_id", "you", True, "Ref → Users", "U01", "Whose records are being shared.", ("ref", "Users")),
            c("shared_with_user_id", "you", False, "Ref → Users", "", "Who gets access. Leave BLANK to share with every family member.", ("ref", "Users")),
            c("section", "you", True, "Enum", "Medicines", "Medicines = prescriptions, doses and the dose log. Care team = doctors, reasons and hospital numbers.", ("list", "Section")),
            c("access", "you", True, "Enum", "Edit", "View = can see. Edit = can see and change. Ticking pills is always the owner's only.", ("list", "Access")),
            c("created_at", "app", False, "DateTime", "", "Filled by the app."),
            c("created_by", "app", False, "Ref → Users", "", "Filled by the app."),
        ],
        "examples": [
            ["SH01", "U01", "", "Medicines", "Edit", "", ""],
            ["SH02", "U01", "", "Care team", "Edit", "", ""],
        ],
    },
    # ---------- Libraries (shared, no person attached) ----------
    "Medicines": {
        "purpose": "The family medicine library. No person and no hospital: who takes it lives in Prescriptions.",
        "cols": [
            key("medicine_id", "MED01", "Unique ID. The app creates one for you when you add from the app."),
            c("generic_name", "you", True, "Text", "Metformin", "Active ingredient, as printed on the box."),
            c("brand_name", "you", False, "Text", "Glucophage", "Brand name, if any."),
            c("strength", "you", False, "Text", "500 mg", "Strength of one tablet, capsule, ml or pen. Different strengths are different library entries."),
            c("form", "you", True, "Enum", "Tablet", "Use Injection for injections.", ("list", "Form")),
            c("purpose", "you", False, "Text", "Blood sugar", "What it is for, in plain words. Visible to every family member."),
            c("notes", "you", False, "Text", "Do not crush", "Facts about the medicine itself. Nothing personal: every family member can read this."),
            c("photo_box", "app", False, "Image (Drive link)", "", "Filled when a photo is added in the app."),
            c("photo_packet_front", "app", False, "Image (Drive link)", "", "Filled by the app."),
            c("photo_packet_back", "app", False, "Image (Drive link)", "", "Filled by the app."),
            c("photo_pill_front", "app", False, "Image (Drive link)", "", "Filled by the app."),
            c("photo_pill_back", "app", False, "Image (Drive link)", "", "Filled by the app."),
        ] + AUDIT,
        "examples": [
            ["MED01", "Metformin", "Glucophage", "500 mg", "Tablet", "Blood sugar", "", "", "", "", "", "", "", "", "", ""],
        ],
    },
    "Hospitals": {
        "purpose": "Hospitals and clinics the family uses. Shared. The HN is personal, so it lives in HospitalNumbers.",
        "cols": [
            key("hospital_id", "HOS01", "Unique ID."),
            c("name", "you", True, "Text", "Riverside General Hospital", "Hospital or clinic name."),
            c("phone", "you", False, "Phone", "02-555-0110", "Main or appointment line."),
            c("address", "you", False, "Text", "", "Optional."),
            c("map_link", "you", False, "URL", "", "A Google Maps link, optional."),
            c("notes", "you", False, "Text", "", "Optional, e.g. parking, which building."),
        ] + AUDIT,
        "examples": [
            ["HOS01", "Riverside General Hospital", "02-555-0110", "", "", "", "", "", "", ""],
        ],
    },
    "Doctors": {
        "purpose": "Doctors the family sees. Shared: one doctor can look after several family members, at several hospitals.",
        "cols": [
            key("doctor_id", "DOC01", "Unique ID."),
            c("name", "you", True, "Text", "Dr. Somchai K.", "Full name."),
            c("specialty", "you", False, "Text", "Cardiology", "The doctor's specialty. Why they see a particular person goes in CareTeam."),
            c("phone", "you", False, "Phone", "02-555-0112", "Clinic or nurse station line."),
            c("other_contact", "you", False, "Text", "", "LINE, email, room number."),
            c("photo", "app", False, "Image (Drive link)", "", "The doctor's photo, filled when added in the app. Readable by anyone holding the link."),
            c("notes", "you", False, "Text", "", "Optional."),
        ] + AUDIT,
        "examples": [
            ["DOC01", "Dr. Somchai K.", "Cardiology", "02-555-0112", "", "", "", "", "", "", ""],
        ],
    },
    "DoctorHospitals": {
        "purpose": "Which hospitals each doctor works at. One row per doctor per hospital.",
        "cols": [
            key("doctor_hospital_id", "DH01", "Unique ID."),
            c("doctor_id", "you", True, "Ref → Doctors", "DOC01", "The doctor.", ("ref", "Doctors")),
            c("hospital_id", "you", True, "Ref → Hospitals", "HOS01", "A hospital they work at.", ("ref", "Hospitals")),
        ],
        "examples": [["DH01", "DOC01", "HOS01"]],
    },
    # ---------- Personal records ----------
    "HospitalNumbers": {
        "purpose": "Each person's HN at each hospital. One row per person per hospital.",
        "cols": [
            key("hn_id", "HN01", "Unique ID."),
            c("user_id", "you", True, "Ref → Users", "U01", "Whose HN.", ("ref", "Users")),
            c("hospital_id", "you", True, "Ref → Hospitals", "HOS01", "At which hospital.", ("ref", "Hospitals")),
            c("hn", "you", True, "Text", "0045821", "The hospital number. Stored as text so leading zeros stay."),
            c("notes", "you", False, "Text", "", "Optional."),
        ] + AUDIT,
        "examples": [["HN01", "U01", "HOS01", "0045821", "", "", "", "", ""]],
    },
    "CareTeam": {
        "purpose": "Which doctor sees which person, where, and why. A doctor can appear for several people with different reasons.",
        "cols": [
            key("care_id", "CT01", "Unique ID."),
            c("user_id", "you", True, "Ref → Users", "U01", "The patient.", ("ref", "Users")),
            c("doctor_id", "you", True, "Ref → Doctors", "DOC01", "The doctor.", ("ref", "Doctors")),
            c("hospital_id", "you", True, "Ref → Hospitals", "HOS01", "Where this doctor sees this person, so the right HN shows.", ("ref", "Hospitals")),
            c("reason", "you", False, "Text", "Blood pressure follow-up", "Why this doctor sees this person."),
            c("active", "you", True, "Yes/No", "TRUE", "FALSE when this person no longer sees this doctor. Kept for history.", ("list", "YesNo")),
        ] + AUDIT,
        "examples": [["CT01", "U01", "DOC01", "HOS01", "Blood pressure follow-up", "TRUE", "", "", "", ""]],
    },
    "Prescriptions": {
        "purpose": "What each person takes NOW. One row per person per medicine. A change updates this row and adds a line to PrescriptionChanges.",
        "cols": [
            key("prescription_id", "RX01", "Unique ID."),
            c("user_id", "you", True, "Ref → Users", "U01", "Who takes it.", ("ref", "Users")),
            c("medicine_id", "you", True, "Ref → Medicines", "MED01", "Which medicine from the library.", ("ref", "Medicines")),
            c("frequency", "you", True, "Enum", "Daily", "Daily, Every N days, Weekdays, or As needed.", ("list", "Frequency")),
            c("every_n_days", "you", False, "Number", "", "Only for Every N days. Every other day = 2, every 3 weeks = 21.", ("number",)),
            c("weekdays", "you", False, "EnumList", "", "Only for Weekdays. Comma separated: Mon, Thu"),
            c("count_from", "you", False, "Date", "", "Only for Every N days: the date the counting starts (a day it is due)."),
            c("meal_timing", "you", True, "Enum", "After meal", "Before meal, After meal, With meal, or Any time.", ("list", "MealTiming")),
            c("doctor_id", "you", False, "Ref → Doctors", "DOC01", "The doctor who prescribed the current dose.", ("ref", "Doctors")),
            c("status", "you", True, "Enum", "Active", "Active or Stopped. Stopped prescriptions are kept for history, never deleted.", ("list", "Status")),
            c("started_on", "you", True, "Date", "2026-01-10", "When this person first started this medicine."),
            c("notes", "you", False, "Text", "", "Personal instructions for this person, e.g. 'skip if BP under 100'."),
        ] + AUDIT,
        "examples": [["RX01", "U01", "MED01", "Daily", "", "", "", "After meal", "DOC01", "Active", "2026-01-10", "", "", "", "", ""]],
    },
    "PrescriptionDoses": {
        "purpose": "How much at each time of day. One row per prescription per time of day (e.g. 1 tablet Morning, 2 tablets Evening).",
        "cols": [
            key("dose_id", "DS01", "Unique ID."),
            c("prescription_id", "you", True, "Ref → Prescriptions", "RX01", "Which prescription.", ("ref", "Prescriptions")),
            c("time_of_day", "you", True, "Enum", "Morning", "Morning, Noon, Evening or Bedtime. As needed prescriptions have no dose rows.", ("list", "TimeOfDay")),
            c("amount", "you", True, "Number", "1", "How many. Use 0.5 for half.", ("number",)),
            c("unit", "you", True, "Enum", "tablet", "Unit of the amount.", ("list", "Unit")),
        ],
        "examples": [
            ["DS01", "RX01", "Morning", "1", "tablet"],
            ["DS02", "RX01", "Evening", "2", "tablet"],
        ],
    },
    "PrescriptionChanges": {
        "purpose": "Append-only history. The app adds one line for every start, change, stop or restart. Never edit or delete rows.",
        "cols": [
            key("change_id", "CH01", "Unique ID."),
            c("prescription_id", "app", True, "Ref → Prescriptions", "RX01", "Which prescription changed.", ("ref", "Prescriptions")),
            c("changed_at", "app", True, "DateTime", "2026-01-10 09:00", "When."),
            c("changed_by", "app", True, "Ref → Users", "U02", "Who made the change in the app.", ("ref", "Users")),
            c("change_type", "app", True, "Enum", "Started", "Started, Dose changed, Schedule changed, Stopped, Restarted, Corrected.", ("list", "ChangeType")),
            c("doctor_id", "app", False, "Ref → Doctors", "DOC01", "Doctor who ordered the change.", ("ref", "Doctors")),
            c("reason", "app", False, "Text", "Started for blood sugar", "Why."),
            c("before", "app", False, "Text", "", "What it was, in words. Blank for Started."),
            c("after", "app", False, "Text", "Every day, after meal: Morning 1 tablet, Evening 2 tablets", "What it became, in words."),
        ],
        "examples": [["CH01", "RX01", "2026-01-10 09:00", "U02", "Started", "DOC01", "Started for blood sugar", "", "Every day, after meal: Morning 1 tablet, Evening 2 tablets"]],
    },
    "DoseLog": {
        "purpose": "One row per tick. Records the amount actually taken, so the record stays true after a dose changes.",
        "cols": [
            key("log_id", "", "Created by the app."),
            c("prescription_id", "app", True, "Ref → Prescriptions", "", "Which prescription.", ("ref", "Prescriptions")),
            c("date", "app", True, "Date", "", "The day the dose was due."),
            c("time_of_day", "app", True, "Enum", "", "Morning, Noon, Evening or Bedtime.", ("list", "TimeOfDay")),
            c("amount_taken", "app", True, "Number", "", "The amount at the moment it was ticked."),
            c("unit", "app", True, "Enum", "", "Unit at the moment it was ticked.", ("list", "Unit")),
            c("status", "app", True, "Enum", "", "Taken or Skipped.", ("list", "DoseStatus")),
            c("taken_at", "app", True, "DateTime", "", "When it was ticked."),
            c("taken_by", "app", True, "Ref → Users", "", "Who ticked it (always the owner).", ("ref", "Users")),
            c("note", "app", False, "Text", "", "Optional, e.g. injection site."),
        ],
        "examples": [],
    },
    "EmergencyCards": {
        "purpose": "One card per person. Readable by the whole family, and without login except the HN. Keep private things out.",
        "cols": [
            c("user_id", "you", True, "Key + Ref → Users", "U01", "One row per person.", ("ref", "Users"), "key"),
            c("full_name", "you", False, "Text", "Somsak (Dad)", "Name as on ID card."),
            c("date_of_birth", "you", False, "Date", "1953-04-12", "YYYY-MM-DD."),
            c("blood_type", "you", False, "Enum", "O+", "", ("list", "BloodType")),
            c("allergies", "you", False, "Text", "Penicillin (rash); Shellfish (hives)", "Separate with ; and put the reaction in brackets."),
            c("conditions", "you", False, "Text", "Hypertension; Diabetes", "Separate with ;"),
            c("contact1_name", "you", False, "Text", "Ooa", ""),
            c("contact1_relation", "you", False, "Text", "Daughter", ""),
            c("contact1_phone", "you", False, "Phone", "081-555-0142", ""),
            c("contact2_name", "you", False, "Text", "", ""),
            c("contact2_relation", "you", False, "Text", "", ""),
            c("contact2_phone", "you", False, "Phone", "", ""),
            c("notes", "you", False, "Text", "", "Things a paramedic must know, e.g. pacemaker. Readable WITHOUT login."),
            c("updated_at", "app", False, "DateTime", "", "Filled by the app."),
            c("updated_by", "app", False, "Ref → Users", "", "Filled by the app."),
        ],
        "examples": [["U01", "Somsak (Dad)", "1953-04-12", "O+", "Penicillin (rash)", "Hypertension; Diabetes", "Ooa", "Daughter", "", "", "", "", "", "", ""]],
    },
    # ---------- App ----------
    "Settings": {
        "purpose": "App settings. Do not rename keys.",
        "cols": [
            c("key", "you", True, "Key (text)", "photo_folder_id", "Setting name.", None, "key"),
            c("value", "you", True, "Text", PHOTO_FOLDER_ID, "Setting value."),
            c("notes", "you", False, "Text", "", "What it is for."),
        ],
        "examples_real": [["photo_folder_id", PHOTO_FOLDER_ID, "Drive folder where the app files medicine and doctor photos"]],
    },
}

# Text-format columns: dates, times, phones, HN, codes, ids, EnumList
TEXT_TYPES = ("Date", "DateTime", "Phone", "Text (6 digits)", "EnumList")
TEXT_NAMES = {"hn", "reset_code", "weekdays"}

MIGRATION = [
    ("Users", "Users", "Copy user_id, display_name, active. Copy password_hash too, so everyone keeps their password. role: Dad → Primary, Family stays Family. Drop line_user_id and the reminder_* columns."),
    ("Sharing", "Sharing", "One row per old rule, with a new sharing_id. section Medications → Medicines (add a second row with section Care team if doctors should be shared too). shared_with ALL → leave blank. Choose View or Edit deliberately: they now mean different things."),
    ("Medications", "Medicines", "med_id → medicine_id. Copy names, strength, form, purpose, notes and the 5 photo links. Drop owner_user_id."),
    ("Hospitals", "Hospitals + HospitalNumbers", "One Hospitals row per real hospital (hospital_name → name, phone, address). For each old row with an hn, add a HospitalNumbers row: owner_user_id → user_id, the hospital_id, the hn."),
    ("Doctors", "Doctors + DoctorHospitals + CareTeam", "One Doctors row per real doctor (doctor_name → name, specialty, phone, other_contact). Old hospital_id → a DoctorHospitals row. Old owner_user_id → a CareTeam row (user_id, doctor_id, hospital_id, reason)."),
    ("Schedules", "Prescriptions + PrescriptionDoses", "Only rows still being taken (end_date blank). One Prescriptions row per person + medicine: frequency_type → frequency, every_n_days, weekdays, start_date → started_on (and count_from for Every N days), meal_timing, prescribed_by → doctor_id, status Active. Then one PrescriptionDoses row for each of morning/noon/evening/bedtime that was TRUE, with dose_amount → amount and dose_unit → unit."),
    ("Schedules (ended rows)", "PrescriptionChanges (optional)", "If you want old history kept, add one line per past change: change_type, reason (change_reason), doctor_id, and before/after in words. Otherwise skip."),
    ("DoseLog", "DoseLog (optional)", "Optional. slot morning → time_of_day Morning; add amount_taken and unit from that day's dose; logged_at → taken_at; logged_by → taken_by."),
    ("Emergency", "EmergencyCards", "Same columns: owner_user_id → user_id. Copy the rest as is."),
    ("Settings", "Settings", "photo_folder_id is already filled in."),
    ("Appointments, Vaccinations, Vitals, Documents", "(not in this version)", "Left out until their phase."),
]

wb = Workbook()
rd = wb.active
rd.title = "README"
lines = [
    ("CyberHealth database, version 2", "title"),
    ("", None),
    ("WHAT THIS IS", "head"),
    ("An AppSheet-style database: every tab is a table, its first column is the key, and *_id columns link tables together.", None),
    ("Medicines, Hospitals and Doctors are shared family libraries. Everything personal points at a person through user_id.", None),
    ("", None),
    ("HOW TO USE IT (do not overwrite the live Sheet)", "head"),
    ("1. Upload this file to Google Drive and open it with Google Sheets (File > Save as Google Sheets). This becomes a NEW Sheet.", None),
    ("2. Keep it private. Only its Apps Script reads it.", None),
    ("3. Grey rows are EXAMPLES. Replace them with your real data or delete them. Settings row 2 is real: keep it.", None),
    ("4. Move your data across using the MOVING YOUR DATA table below.", None),
    ("5. Keep the OLD Sheet running the live app until the new app is deployed on this Sheet.", None),
    ("", None),
    ("RULES", "head"),
    ("Do not rename tabs or headers: the app finds columns by name.", None),
    ("Header colours: PURPLE = key, RED = required, BLUE = optional, GREY = filled by the app.", None),
    ("IDs only need to be unique within their tab. IDs you type by hand (U01, MED01) are fine; the app makes its own for new rows.", None),
    ("Dates are YYYY-MM-DD. Times of day are Morning, Noon, Evening, Bedtime.", None),
    ("Never delete rows from Prescriptions, PrescriptionChanges or DoseLog: stop a prescription instead.", None),
    ("Photos (medicines and doctors) are readable by anyone holding the photo's link. The folder stays private.", None),
    ("The emergency card (except HN) can be opened without logging in. Keep private things out of it.", None),
    ("", None),
    ("MOVING YOUR DATA", "head"),
]
for i, (text, style) in enumerate(lines, start=1):
    cell = rd.cell(row=i, column=1, value=text)
    if style == "title":
        cell.font = Font(bold=True, size=15, color="4B2E83")
    elif style == "head":
        cell.font = Font(bold=True, color="4B2E83")
row = len(lines) + 1
for j, h in enumerate(["Old tab", "New tab(s)", "How"], start=1):
    cell = rd.cell(row=row, column=j, value=h); cell.font = WHITE; cell.fill = FILL_OPT
for old, new, how in MIGRATION:
    row += 1
    rd.cell(row=row, column=1, value=old)
    rd.cell(row=row, column=2, value=new)
    rd.cell(row=row, column=3, value=how).alignment = Alignment(wrap_text=True, vertical="top")
row += 2
rd.cell(row=row, column=1, value="COLUMN GUIDE").font = Font(bold=True, color="4B2E83")
row += 1
guide_headers = ["Tab", "Column", "Key/Required", "Filled by", "Type", "Example", "Description"]
for j, h in enumerate(guide_headers, start=1):
    cell = rd.cell(row=row, column=j, value=h); cell.font = WHITE; cell.fill = FILL_OPT
for tab, spec in TABLES.items():
    row += 1
    rd.cell(row=row, column=1, value=tab).font = Font(bold=True)
    rd.cell(row=row, column=7, value=spec["purpose"]).font = Font(italic=True)
    for name, kind, who, req, typ, ex, desc, _ in spec["cols"]:
        row += 1
        values = [tab, name, "Key" if kind == "key" else ("Yes" if req else ""), "App" if who == "app" else "You", typ, ex, desc]
        for j, v in enumerate(values, start=1):
            rd.cell(row=row, column=j, value=v)
for col, w in zip("ABCDEFG", [22, 24, 12, 10, 22, 30, 90]):
    rd.column_dimensions[col].width = w

ls = wb.create_sheet("Lists")
list_ranges = {}
for j, (name, values) in enumerate(LISTS.items(), start=1):
    letter = get_column_letter(j)
    cell = ls.cell(row=1, column=j, value=name); cell.font = WHITE; cell.fill = FILL_OPT
    for i, v in enumerate(values, start=2):
        ls.cell(row=i, column=j, value=v)
    list_ranges[name] = f"Lists!${letter}$2:${letter}${len(values) + 1}"
    ls.column_dimensions[letter].width = 18

for tab, spec in TABLES.items():
    ws = wb.create_sheet(tab)
    cols = spec["cols"]
    for j, (name, kind, who, req, typ, ex, desc, val) in enumerate(cols, start=1):
        letter = get_column_letter(j)
        cell = ws.cell(row=1, column=j, value=name)
        cell.font = WHITE
        cell.fill = FILL_KEY if kind == "key" else FILL_APP if who == "app" else FILL_REQ if req else FILL_OPT
        note = f"{'KEY' if kind == 'key' else 'REQUIRED' if req else 'Optional'} | {'Filled by app' if who == 'app' else 'You fill'}\nType: {typ}"
        if desc:
            note += f"\n{desc}"
        cell.comment = Comment(note, "CyberHealth", width=280, height=120)
        ws.column_dimensions[letter].width = max(14, min(34, len(name) + 6))
        text = typ.startswith(TEXT_TYPES) or name in TEXT_NAMES or name.endswith("_id") or kind == "key"
        for r in range(2, MAX_ROW + 1):
            if text:
                ws[f"{letter}{r}"].number_format = "@"
        if val:
            rng = f"{letter}2:{letter}{MAX_ROW}"
            if val[0] == "list":
                dv = DataValidation(type="list", formula1=f"={list_ranges[val[1]]}", allow_blank=True)
            elif val[0] == "ref":
                ref_tab = val[1]
                dv = DataValidation(type="list", formula1=f"={ref_tab}!$A$2:$A${MAX_ROW}", allow_blank=True)
            else:
                dv = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True)
                dv.error, dv.errorTitle = "Please enter a number.", "Number only"
            ws.add_data_validation(dv)
            dv.add(rng)
    real = spec.get("examples_real", [])
    for i, values in enumerate(real + spec.get("examples", []), start=2):
        for j, v in enumerate(values, start=1):
            cell = ws.cell(row=i, column=j, value=v if v != "" else None)
            if i - 2 >= len(real):
                cell.font = EXAMPLE
    ws.freeze_panes = "B2"
    ws.row_dimensions[1].height = 22

wb.move_sheet("Lists", offset=len(wb.sheetnames))
wb.save(OUT)
print("saved", os.path.abspath(OUT))
print("tabs:", wb.sheetnames)
